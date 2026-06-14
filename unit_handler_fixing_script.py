from dataclasses import dataclass, field
from emtp_utils import *
from com_client import *
from unit_models import *

from typing import List
from pathlib import Path
import openpyxl


@dataclass
class UnitManager:

    pv_units: List[Photovoltaic] = field(default_factory=list)
    wt_units: List[WindTurbine] = field(default_factory=list)
    bess_units: List[Bess] = field(default_factory=list)
    synchronous_units: List[Synchronous] = field(default_factory=list)


class UnitExtractor:

    def __init__(self, emtp_object):
        self.emtp_object = emtp_object

        self.units = UnitManager()

    def _get_devices(self):

        cct = self.emtp_object.currentCircuit
        self.devices = cct.devices

    def _classify_devices(self):

        _devices = self.devices

        for device in _devices:
            if "WECC PV" in device.getAttribute(
                "LibType" or "PV_TEMPLATE" in device.getAttribute("LibType")
            ) and "PFV_" in device.getAttribute("Name"):

                unit = Photovoltaic(object=None, unit_object=device)

                self.units.pv_units.append(unit)

            if (
                "WECC W" in device.getAttribute("LibType")
                or "WT_TEMPLATE" in device.getAttribute("LibType")
            ) and "PE_" in device.getAttribute("Name"):
                unit = WindTurbine(object=None, unit_object=device)
                self.units.wt_units.append(unit)

            # elif device.getAttribute("LibType") == "Bess":
            #     self.units.bess_units.append(device)
            elif any(
                prefix in device.getAttribute("Name")[:4]
                for prefix in ["HE_", "HP_", "TER_"]
            ):
                parent_device = device
                subcct = parent_device.subCircuit

                syn_unit = None
                lf_unit = None

                for sub_device in subcct.devices:
                    if "synchronous" in sub_device.getAttribute("LibType").lower():
                        syn_unit = sub_device

                    elif sub_device.getAttribute("LibType") == "Load-Flow Bus":
                        lf_unit = sub_device

                if syn_unit and lf_unit:
                    unit = Synchronous(
                        object=parent_device,
                        unit_object=syn_unit,
                        lf_object=lf_unit,
                    )
                    self.units.synchronous_units.append(unit)

                else:
                    print(
                        f"Error: {device.name} does not have a synchronous or load-flow unit."
                    )

    def execute(self):

        self._get_devices()
        self._classify_devices()


if __name__ == "__main__":

    emtp_client = EmtpComClient(attach_existing=True)
    emtp_object = emtp_client.emtp_object

    if not emtp_object.currentDesign:
        Design.open_design(emtp_object=emtp_object)

    unit_extractor = UnitExtractor(emtp_object=emtp_object)
    unit_extractor.execute()

    # Export to a excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ── PV & WT (same structure) ──────────────────────────────────────
    RES_COLUMNS = [
        "Name",
        "LibType",
        "S (MVA)",
        "Exclude",
        "P (MW)",
        "Q (MVAR)",
        "V_POI",
        "SubmoduleID",
        "posX",
        "posY",
        "Orientation",
    ]

    for sheet_name, unit_list in [
        ("PV", unit_extractor.units.pv_units),
        ("WT", unit_extractor.units.wt_units),
    ]:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(RES_COLUMNS)
        for unit in unit_list:
            attr_machine = Device.get_params_to_dict(
                emtp_object=emtp_object, device=unit.unit_object
            )
            ngen = float(attr_machine.get("Ngen", 0))
            sgen = float(attr_machine.get("Sgen", 0))
            ngen_in_service = float(attr_machine.get("Ngen_in_service", ngen))
            pref_poi = float(attr_machine.get("Pref_poi", 0))
            qpoi_pu = float(attr_machine.get("Qpoi_pu", 0))
            sgn = round(ngen * sgen, 2)
            p = round(pref_poi * sgen * ngen_in_service, 2)
            q = round(qpoi_pu * sgen * ngen_in_service, 2)

            ws.append(
                [
                    unit.unit_object.getAttribute("Name"),
                    unit.unit_object.getAttribute("LibType"),
                    sgn,
                    unit.unit_object.getAttribute("Exclude"),
                    p,
                    q,
                    attr_machine.get("Vpoi_kVRMSLL", ""),
                    unit.unit_object.getAttribute("SubmoduleID"),
                    unit.unit_object.posX,
                    unit.unit_object.posY,
                    unit.unit_object.orientation,
                ]
            )

    # ── Synchronous (different attributes, different objects) ─────────
    SYN_COLUMNS = [
        "Name",
        "LibType",
        "S (MVA)",
        "Exclude",
        "P (MW)",
        "PV/PQ node",
        "V_Setpoint",
        "V_Nominal",
        "V (pu)",
        "Q (MVAr)",
    ]

    ws_syn = wb.create_sheet(title="Synchronous")
    ws_syn.append(SYN_COLUMNS)

    for unit in unit_extractor.units.synchronous_units:

        # /// Ensure script

        try:
            script_devobj = unit.unit_object.getAttribute("Script.DevObj")

            if script_devobj == "":
                lib_type = unit.unit_object.getAttribute("LibType")

                if lib_type in ("Synchronous", "synchronouss machine"):
                    unit.unit_object.setAttribute("Script.DevObj", "machine_sm_d.dwj")
        except:
            print(f"Error setting Script.DevObj for {unit.object.name}")

        try:
            script_devobj = unit.lf_object.getAttribute("Script.DevObj")

            if script_devobj == "":
                lib_type = unit.lf_object.getAttribute("LibType")

                if lib_type == "Load-Flow Bus":
                    unit.lf_object.setAttribute("Script.DevObj", "load_flow_bus_d.dwj")
        except:
            print(f"Error setting Script.DevObj for {unit.object.name}")

        # // Ensure script
        try:
            machine_path = f"{unit.object.name}/{unit.unit_object.name}"
            lf_path = f"{unit.object.name}/{unit.lf_object.name}"

        except:
            print(f"Error determining paths for {unit.object.name}")

        # print("Machine path:", machine_path)
        # print("LF path:", lf_path)
        try:
            attr_lf = Utils.get_params_to_dict_by_path(
                emtp_object=emtp_object,
                device_path=lf_path,
            )
        except:
            attr_lf = None
            print(f"Error getting Load-Flow attributes for {unit.object.name}")

        try:
            attr_machine = Utils.get_params_to_dict_by_path(
                emtp_object=emtp_object,
                device_path=machine_path,
            )
        except:
            attr_machine = None
            print(f"Error getting Machine attributes for {unit.object.name}")

        try:
            if (
                unit.object.getAttribute("Exclude") == ""
                and unit.unit_object.getAttribute("Exclude") == ""
            ):
                in_service = ""
            else:
                in_service = "Ex"
        except:
            print(f"Error determining in_service for {unit.object.name}")

        if attr_machine and attr_lf:
            ws_syn.append(
                [
                    unit.unit_object.getAttribute("Name"),
                    unit.unit_object.getAttribute("LibType"),
                    attr_machine.get("Rating_S", ""),
                    in_service,
                    attr_lf.get("P_set", ""),
                    attr_lf.get("Bus_Type", ""),
                    attr_lf.get("Voltage_Slack", ""),
                    attr_machine.get("Rating_V", ""),
                    round(
                        float(attr_lf.get("Voltage_Slack", ""))
                        / float(attr_machine.get("Rating_V", "")),
                        2,
                    ),
                    attr_lf.get("Q_set", ""),
                ]
            )

        else:
            ws_syn.append(
                [
                    unit.unit_object.getAttribute("Name"),
                    unit.unit_object.getAttribute("LibType"),
                    "",
                    in_service,
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )

    output_excel = Path(__file__).parent / "units_export.xlsx"
    wb.save(output_excel)
    print(f"✅ Excel saved to: {output_excel}")
