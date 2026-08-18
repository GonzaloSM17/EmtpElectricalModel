"""
Rename EMTP generation units using an Excel mapping.

Rules:
1. If LibType contains "TEMPLATE" and the current name exists in Excel:
   - Rename the main-circuit mask.
   - Rename the main internal unit.
   - For SG templates, also rename LF, transformer and SSAA.

2. If LibType contains "TEMPLATE" but the name does not exist in Excel:
   - Do not modify the unit.
   - Add it to the report.

3. If the name exists in Excel but LibType does not contain "TEMPLATE":
   - Do not modify the unit.
   - Add it to the report.

The script generates an Excel report with the complete result.
"""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

sys.path.append(str(Path(__file__).resolve().parent))

from com_client import EmtpComClient
from emtp_utils import Design


# =============================================================================
# USER CONFIGURATION
# =============================================================================

EXCEL_PATH = (
    r"C:\Users\kelly.allendes\Downloads\Unit_Names.xlsx"
)

RESULT_PATH = (
    r"C:\Users\kelly.allendes\Downloads\Unit_Names_Result.xlsx"
)

SHEET_NAME = "Names"

# Excel column containing the current EMTP name.
EXCEL_OLD_NAME_COLUMN = "Old_Name"

# Excel column containing the new name.
EXCEL_NEW_NAME_COLUMN = "New_Name"


# =============================================================================
# GENERAL FUNCTIONS
# =============================================================================

def normalize_name(value) -> str:
    """Convert a value into a clean comparable string."""

    if value is None:
        return ""

    return str(value).strip()


def get_attribute_safe(device, attribute_name: str, default=""):
    """Read an EMTP attribute without stopping the script."""

    try:
        value = device.getAttribute(attribute_name)

        if value is None:
            return default

        return value

    except Exception:
        return default


def get_device_name(device) -> str:
    """Return the device name safely."""

    try:
        return normalize_name(device.name)

    except Exception:
        return normalize_name(
            get_attribute_safe(device, "Name")
        )


def is_template(device) -> bool:
    """Return True when LibType contains TEMPLATE."""

    lib_type = normalize_name(
        get_attribute_safe(device, "LibType")
    )

    return "TEMPLATE" in lib_type.upper()


# =============================================================================
# EXCEL READING
# =============================================================================

def read_name_mapping(
    excel_path: str,
    sheet_name: str,
    old_name_column: str,
    new_name_column: str,
) -> dict[str, str]:
    """Read the old-name/new-name mapping from Excel."""

    workbook = openpyxl.load_workbook(
        excel_path,
        data_only=True,
    )

    if sheet_name not in workbook.sheetnames:
        raise ValueError(
            f"The sheet '{sheet_name}' does not exist. "
            f"Available sheets: {workbook.sheetnames}"
        )

    worksheet = workbook[sheet_name]

    headers = {}

    for column_index, cell in enumerate(
        worksheet[1],
        start=1,
    ):
        header = normalize_name(cell.value)

        if header:
            headers[header] = column_index

    if old_name_column not in headers:
        raise ValueError(
            f"Column '{old_name_column}' was not found "
            f"in sheet '{sheet_name}'."
        )

    if new_name_column not in headers:
        raise ValueError(
            f"Column '{new_name_column}' was not found "
            f"in sheet '{sheet_name}'."
        )

    old_column_index = headers[old_name_column]
    new_column_index = headers[new_name_column]

    mapping = {}

    for row_index in range(
        2,
        worksheet.max_row + 1,
    ):
        old_name = normalize_name(
            worksheet.cell(
                row=row_index,
                column=old_column_index,
            ).value
        )

        new_name = normalize_name(
            worksheet.cell(
                row=row_index,
                column=new_column_index,
            ).value
        )

        if not old_name:
            continue

        if not new_name:
            continue

        if old_name in mapping:
            print(
                f"Warning: duplicated Excel name '{old_name}'. "
                f"The last value will be used."
            )

        mapping[old_name] = new_name

    workbook.close()

    return mapping


# =============================================================================
# UNIT CLASSIFICATION
# =============================================================================

def get_unit_type(name: str, lib_type: str = "") -> str:
    """Classify the unit using its name and LibType."""

    name_upper = normalize_name(name).upper()
    lib_type_upper = normalize_name(lib_type).upper()

    if name_upper.startswith(
        (
            "TER_",
            "GEO_",
            "HE_",
            "HP_",
            "CONDENSADOR_",
        )
    ):
        return "SG"

    if name_upper.startswith("PFV_"):
        return "PV"

    if name_upper.startswith("PMGD_"):
        return "PMGD"

    if name_upper.startswith(
        (
            "PE_",
            "WP_",
            "WT_",
        )
    ):
        return "WIND"

    if name_upper.startswith("BESS_"):
        return "BESS"

    # Fallback classification using LibType.
    if "SG_" in lib_type_upper:
        return "SG"

    if "PV_TEMPLATE" in lib_type_upper:
        return "PV"

    if "DER_TEMPLATE" in lib_type_upper:
        return "PMGD"

    if (
        "WT_TEMPLATE" in lib_type_upper
        or "WIND_TEMPLATE" in lib_type_upper
    ):
        return "WIND"

    if "BESS_TEMPLATE" in lib_type_upper:
        return "BESS"

    return "UNKNOWN"


# =============================================================================
# INTERNAL DEVICE SEARCH
# =============================================================================

def find_internal_devices(mask, unit_type: str) -> dict:
    """
    Find the internal devices that should be renamed.

    For all unit types:
    - main_unit

    For SG:
    - load_flow
    - transformer
    - load
    """

    result = {
        "main_unit": None,
        "load_flow": None,
        "transformer": None,
        "load": None,
    }

    try:
        internal_devices = mask.subCircuit.devices

    except Exception:
        return result

    internal_devices_list = []

    for internal_device in internal_devices:
        internal_devices_list.append(internal_device)

        name = get_device_name(internal_device)

        part = normalize_name(
            get_attribute_safe(
                internal_device,
                "Part",
            )
        )

        lib_type = normalize_name(
            get_attribute_safe(
                internal_device,
                "LibType",
            )
        )

        name_upper = name.upper()
        part_upper = part.upper()
        lib_type_upper = lib_type.upper()

        # ---------------------------------------------------------------------
        # Synchronous generation
        # ---------------------------------------------------------------------
        if unit_type == "SG":

            if (
                part_upper == "SM"
                or lib_type_upper == "SYNCHRONOUS"
                or "SYNCHRONOUS" in lib_type_upper
            ):
                result["main_unit"] = internal_device

            if (
                lib_type_upper == "LOAD-FLOW BUS"
                or name_upper.startswith("LF_")
            ):
                result["load_flow"] = internal_device

            if (
                name_upper.startswith("TR_")
                or "TRANSFORMER" in lib_type_upper
                or part_upper
                in (
                    "YGD_P30",
                    "YD_P30",
                    "DY_M30",
                    "YY",
                    "YGYG",
                    "YGYGD",
                )
            ):
                result["transformer"] = internal_device

            if (
                part_upper == "PQLOAD"
                or name_upper.startswith("SSAA_")
            ):
                result["load"] = internal_device

        # ---------------------------------------------------------------------
        # Photovoltaic and PMGD
        # ---------------------------------------------------------------------
        elif unit_type in ("PV", "PMGD"):

            if (
                "WECC PV" in lib_type_upper
                or "PV" in lib_type_upper
                or "DER" in lib_type_upper
                or part_upper in ("PV", "DER")
            ):
                result["main_unit"] = internal_device

        # ---------------------------------------------------------------------
        # Wind generation
        # ---------------------------------------------------------------------
        elif unit_type == "WIND":

            if (
                "WECC W" in lib_type_upper
                or "WIND" in lib_type_upper
                or "WT" in lib_type_upper
                or part_upper in ("WT", "WIND")
            ):
                result["main_unit"] = internal_device

        # ---------------------------------------------------------------------
        # BESS
        # ---------------------------------------------------------------------
        elif unit_type == "BESS":

            if (
                "AC-DC CONVERTER" in lib_type_upper
                or "BESS" in lib_type_upper
                or "BATTERY" in lib_type_upper
                or part_upper in ("BESS", "VSC")
            ):
                result["main_unit"] = internal_device

    # For PV, PMGD, wind and BESS, the main model is commonly the
    # first or only internal device.
    if (
        unit_type != "SG"
        and result["main_unit"] is None
        and internal_devices_list
    ):
        if len(internal_devices_list) == 1:
            result["main_unit"] = internal_devices_list[0]

        else:
            # Fallback: select the first device that is not an obvious
            # auxiliary object.
            for internal_device in internal_devices_list:
                name = get_device_name(internal_device).upper()

                lib_type = normalize_name(
                    get_attribute_safe(
                        internal_device,
                        "LibType",
                    )
                ).upper()

                if (
                    "LOAD-FLOW BUS" in lib_type
                    or name.startswith("LF_")
                    or name.startswith("TR_")
                    or name.startswith("SSAA_")
                ):
                    continue

                result["main_unit"] = internal_device
                break

    return result


# =============================================================================
# RENAMING FUNCTIONS
# =============================================================================

def rename_device(device, new_name: str) -> tuple[bool, str]:
    """Rename an EMTP device."""

    if device is None:
        return False, "Device not found"

    try:
        device.setAttribute("Name", new_name)

        resulting_name = get_device_name(device)

        if resulting_name != new_name:
            return (
                False,
                f"Expected '{new_name}', "
                f"but the resulting name is '{resulting_name}'",
            )

        return True, ""

    except Exception as error:
        return False, str(error)


def rename_unit(
    mask,
    old_name: str,
    new_name: str,
    unit_type: str,
) -> dict:
    """Rename a template mask and its internal devices."""

    result = {
        "old_name": old_name,
        "new_name": new_name,
        "unit_type": unit_type,
        "mask": False,
        "main_unit": False,
        "load_flow": True,
        "transformer": True,
        "load": True,
        "errors": [],
    }

    # Locate internal devices before renaming the mask.
    internal_devices = find_internal_devices(
        mask=mask,
        unit_type=unit_type,
    )

    # -------------------------------------------------------------------------
    # Rename main-circuit mask
    # -------------------------------------------------------------------------
    success, error = rename_device(
        mask,
        new_name,
    )

    result["mask"] = success

    if error:
        result["errors"].append(
            f"Mask: {error}"
        )

    # -------------------------------------------------------------------------
    # Rename main internal model
    # -------------------------------------------------------------------------
    success, error = rename_device(
        internal_devices["main_unit"],
        new_name,
    )

    result["main_unit"] = success

    if error:
        result["errors"].append(
            f"Internal unit: {error}"
        )

    # -------------------------------------------------------------------------
    # Additional SG internal devices
    # -------------------------------------------------------------------------
    if unit_type == "SG":

        success, error = rename_device(
            internal_devices["load_flow"],
            f"LF_{new_name}",
        )

        result["load_flow"] = success

        if error:
            result["errors"].append(
                f"Load-flow: {error}"
            )

        success, error = rename_device(
            internal_devices["transformer"],
            f"TR_{new_name}",
        )

        result["transformer"] = success

        if error:
            result["errors"].append(
                f"Transformer: {error}"
            )

        if internal_devices["load"] is not None:

            success, error = rename_device(
                internal_devices["load"],
                f"SSAA_{new_name}",
            )

            result["load"] = success

            if error:
                result["errors"].append(
                    f"Auxiliary load: {error}"
                )

        else:
            # A missing auxiliary load is allowed.
            result["load"] = True

    return result


# =============================================================================
# REPORT FUNCTIONS
# =============================================================================

def apply_report_format(worksheet):
    """Apply basic formatting to a result worksheet."""

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="A7D0F5",
    )

    header_font = Font(
        bold=True,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        maximum_length = 0

        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            maximum_length = max(
                maximum_length,
                len(value),
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            maximum_length + 3,
            70,
        )


def create_result_excel(
    result_path: str,
    rename_results: list[dict],
    templates_without_excel: list[dict],
    excel_non_templates: list[dict],
    excel_not_found: list[dict],
):
    """Create the final Excel report."""

    workbook = openpyxl.Workbook()

    # -------------------------------------------------------------------------
    # Renamed templates
    # -------------------------------------------------------------------------
    renamed_sheet = workbook.active
    renamed_sheet.title = "Renamed_Templates"

    renamed_sheet.append(
        [
            "Old_Name",
            "New_Name",
            "Unit_Type",
            "Status",
            "Mask_Renamed",
            "Internal_Unit_Renamed",
            "LF_Renamed",
            "Transformer_Renamed",
            "SSAA_Renamed",
            "Details",
        ]
    )

    for result in rename_results:
        renamed_sheet.append(
            [
                result["old_name"],
                result["new_name"],
                result["unit_type"],
                result["status"],
                result["mask"],
                result["main_unit"],
                result["load_flow"],
                result["transformer"],
                result["load"],
                result["details"],
            ]
        )

    # -------------------------------------------------------------------------
    # Templates not included in Excel
    # -------------------------------------------------------------------------
    templates_sheet = workbook.create_sheet(
        "Templates_Not_In_Excel"
    )

    templates_sheet.append(
        [
            "EMTP_Name",
            "LibType",
            "Status",
        ]
    )

    for item in templates_without_excel:
        templates_sheet.append(
            [
                item["name"],
                item["lib_type"],
                item["status"],
            ]
        )

    # -------------------------------------------------------------------------
    # Excel names found in EMTP but not templates
    # -------------------------------------------------------------------------
    non_template_sheet = workbook.create_sheet(
        "Excel_Non_Templates"
    )

    non_template_sheet.append(
        [
            "Current_Name",
            "Requested_New_Name",
            "LibType",
            "Status",
        ]
    )

    for item in excel_non_templates:
        non_template_sheet.append(
            [
                item["old_name"],
                item["new_name"],
                item["lib_type"],
                item["status"],
            ]
        )

    # -------------------------------------------------------------------------
    # Excel names not found anywhere in EMTP
    # -------------------------------------------------------------------------
    not_found_sheet = workbook.create_sheet(
        "Excel_Not_Found"
    )

    not_found_sheet.append(
        [
            "Old_Name",
            "New_Name",
            "Status",
        ]
    )

    for item in excel_not_found:
        not_found_sheet.append(
            [
                item["old_name"],
                item["new_name"],
                item["status"],
            ]
        )

    for worksheet in workbook.worksheets:
        apply_report_format(worksheet)

    workbook.save(result_path)


# =============================================================================
# MAIN
# =============================================================================

if __name__ == "__main__":

    # -------------------------------------------------------------------------
    # Read Excel mapping
    # -------------------------------------------------------------------------
    name_mapping = read_name_mapping(
        excel_path=EXCEL_PATH,
        sheet_name=SHEET_NAME,
        old_name_column=EXCEL_OLD_NAME_COLUMN,
        new_name_column=EXCEL_NEW_NAME_COLUMN,
    )

    # -------------------------------------------------------------------------
    # Connect to EMTP
    # -------------------------------------------------------------------------
    emtp_client = EmtpComClient(
        attach_existing=True
    )

    emtp_object = emtp_client.emtp_object

    if not emtp_object.currentDesign:
        Design.open_design(
            emtp_object=emtp_object
        )

    circuit = emtp_object.currentCircuit

    # Convert COM collection into a Python list before renaming devices.
    main_devices = [
        device
        for device in circuit.devices
    ]

    rename_results = []
    templates_without_excel = []
    excel_non_templates = []

    # Excel names that were found in EMTP, regardless of whether they
    # were templates.
    found_excel_names = set()

    # -------------------------------------------------------------------------
    # Process all devices in the main circuit
    # -------------------------------------------------------------------------
    for mask in main_devices:

        old_name = get_device_name(mask)

        lib_type = normalize_name(
            get_attribute_safe(
                mask,
                "LibType",
            )
        )

        template_status = is_template(mask)
        exists_in_excel = old_name in name_mapping

        # ---------------------------------------------------------------------
        # Case 1:
        # TEMPLATE + name exists in Excel -> rename.
        # ---------------------------------------------------------------------
        if template_status and exists_in_excel:

            found_excel_names.add(old_name)

            new_name = name_mapping[old_name]

            unit_type = get_unit_type(
                name=old_name,
                lib_type=lib_type,
            )

            rename_result = rename_unit(
                mask=mask,
                old_name=old_name,
                new_name=new_name,
                unit_type=unit_type,
            )

            required_results = [
                rename_result["mask"],
                rename_result["main_unit"],
            ]

            if unit_type == "SG":
                required_results.extend(
                    [
                        rename_result["load_flow"],
                        rename_result["transformer"],
                        rename_result["load"],
                    ]
                )

            if all(required_results):
                status = "Renamed successfully"

            elif any(required_results):
                status = "Partially renamed"

            else:
                status = "Rename failed"

            rename_result["status"] = status

            rename_result["details"] = " | ".join(
                rename_result["errors"]
            )

            rename_results.append(
                rename_result
            )

            print(
                f"{old_name} -> {new_name}: {status}"
            )

            continue

        # ---------------------------------------------------------------------
        # Case 2:
        # TEMPLATE + name is not in Excel -> do not modify; report it.
        # ---------------------------------------------------------------------
        if template_status and not exists_in_excel:

            templates_without_excel.append(
                {
                    "name": old_name,
                    "lib_type": lib_type,
                    "status": (
                        "Template not found in Excel. "
                        "No name change was applied."
                    ),
                }
            )

            continue

        # ---------------------------------------------------------------------
        # Case 3:
        # Name exists in Excel but it is not a TEMPLATE.
        # Do not modify; report it.
        # ---------------------------------------------------------------------
        if not template_status and exists_in_excel:

            found_excel_names.add(old_name)

            excel_non_templates.append(
                {
                    "old_name": old_name,
                    "new_name": name_mapping[old_name],
                    "lib_type": lib_type,
                    "status": (
                        "Name found in Excel, but LibType "
                        "does not contain TEMPLATE. "
                        "No name change was applied."
                    ),
                }
            )

            continue

    # -------------------------------------------------------------------------
    # Excel names that do not exist anywhere in the main EMTP circuit
    # -------------------------------------------------------------------------
    excel_not_found = []

    for old_name, new_name in name_mapping.items():

        if old_name not in found_excel_names:

            excel_not_found.append(
                {
                    "old_name": old_name,
                    "new_name": new_name,
                    "status": (
                        "Excel name was not found in "
                        "the main EMTP circuit."
                    ),
                }
            )

    # -------------------------------------------------------------------------
    # Save EMTP design
    # -------------------------------------------------------------------------
    Design.save(
        emtp_object=emtp_object
    )

    # -------------------------------------------------------------------------
    # Create result report
    # -------------------------------------------------------------------------
    create_result_excel(
        result_path=RESULT_PATH,
        rename_results=rename_results,
        templates_without_excel=templates_without_excel,
        excel_non_templates=excel_non_templates,
        excel_not_found=excel_not_found,
    )

    # -------------------------------------------------------------------------
    # Final summary
    # -------------------------------------------------------------------------
    print("")
    print("=" * 70)
    print("PROCESS COMPLETED")
    print("=" * 70)

    print(
        f"Templates renamed: "
        f"{len(rename_results)}"
    )

    print(
        f"Templates not found in Excel: "
        f"{len(templates_without_excel)}"
    )

    print(
        f"Excel names corresponding to non-template devices: "
        f"{len(excel_non_templates)}"
    )

    print(
        f"Excel names not found in EMTP: "
        f"{len(excel_not_found)}"
    )

    print(
        f"Result file: {RESULT_PATH}"
    )